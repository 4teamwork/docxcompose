"""Compare append performance and package contents against a local git revision.

Run after an editable install: python benchmarks/styles.py --baseline-ref origin/master
"""

import argparse
import gc
import hashlib
import json
import platform
import random
import statistics
import subprocess
import time
from importlib.metadata import version
from io import BytesIO

from docx import Document
from docx.enum.style import WD_STYLE_TYPE

from docxcompose import composer


def document_bytes(document):
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def run(composer_class, master_bytes, child_bytes, preserve_styles):
    master = Document(BytesIO(master_bytes))
    children = [Document(BytesIO(child_bytes)) for _ in range(8)]
    instance = composer_class(master, preserve_styles=preserve_styles)
    random.seed(42)
    started = time.perf_counter()
    for child in children:
        instance.append(child)
    elapsed = time.perf_counter() - started
    parts = {
        str(part.partname): hashlib.sha256(part.blob).hexdigest()
        for part in master.part.package.parts
    }
    return elapsed, parts


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--baseline-ref", default="origin/master")
    args = parser.parse_args()
    baseline_ref = subprocess.check_output(
        ["git", "rev-parse", args.baseline_ref], text=True
    ).strip()
    source = subprocess.check_output(
        ["git", "show", f"{baseline_ref}:docxcompose/composer.py"], text=True
    )
    namespace = {"__file__": composer.__file__, "__name__": "benchmark_baseline"}
    exec(compile(source, "baseline/composer.py", "exec"), namespace)
    classes = {"baseline": namespace["Composer"], "candidate": composer.Composer}
    results = []
    for unused_styles, preserve_styles in ((0, False), (500, False), (500, True)):
        master = Document()
        for index in range(unused_styles):
            master.styles.add_style(f"Unused{index}", WD_STYLE_TYPE.PARAGRAPH)
        child = Document()
        child.styles.add_style("Imported", WD_STYLE_TYPE.PARAGRAPH)
        child.styles["Heading 1"].font.bold = False
        for index in range(50):
            child.add_paragraph(
                f"Paragraph {index}", style="Imported" if index % 2 else "Heading 1"
            )
        master_bytes, child_bytes = document_bytes(master), document_bytes(child)
        samples = {name: [] for name in classes}
        expected = None
        for repetition in range(6):
            order = list(classes) if repetition % 2 == 0 else list(reversed(classes))
            for name in order:
                gc.collect()
                elapsed, parts = run(
                    classes[name], master_bytes, child_bytes, preserve_styles
                )
                if expected is None:
                    expected = parts
                assert parts == expected, (unused_styles, preserve_styles, name)
                if repetition:
                    samples[name].append(elapsed)
        results.append(
            {
                "unused_master_styles": unused_styles,
                "preserve_styles": preserve_styles,
                "children": 8,
                "paragraphs_per_child": 50,
                "seconds": samples,
                "median_seconds": {
                    name: statistics.median(values) for name, values in samples.items()
                },
                "package_parts_equal": True,
            }
        )
    print(
        json.dumps(
            {
                "baseline_ref": baseline_ref,
                "python": platform.python_version(),
                "platform": platform.system(),
                "architecture": platform.machine(),
                "dependencies": {
                    name: version(name) for name in ("python-docx", "lxml", "babel")
                },
                "timing": "append only; input loading and package comparison excluded",
                "results": results,
            },
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
